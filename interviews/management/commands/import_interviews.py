from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import IntegrityError, transaction

from interviews.models import Interview, User

REQUIRED_COLUMNS = ('ExternalID', 'Title', 'IntervieweeEmail', 'PanelistEmails')


class Command(BaseCommand):
    help = (
        'Import (upsert) Interview rows and their panels from an xlsx. '
        'Interviews are created unscheduled; the solver assigns timeslots later. '
        'PanelistEmails is semicolon-separated. Referenced users must already exist.'
    )

    def add_arguments(self, parser):
        parser.add_argument('path', type=Path, help='Path to the source xlsx file')
        parser.add_argument('--sheet', default='Interviews', help='Worksheet name (default: Interviews)')
        parser.add_argument('--dry-run', action='store_true', help='Parse and validate but do not write')

    def handle(self, *args, path, sheet, dry_run, **_):
        import openpyxl

        if not path.exists():
            raise CommandError(f'file not found: {path}')

        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        if sheet not in wb.sheetnames:
            raise CommandError(f'sheet {sheet!r} not found; available: {wb.sheetnames}')

        rows_iter = wb[sheet].iter_rows(values_only=True)
        headers = next(rows_iter, None)
        if not headers:
            raise CommandError('sheet is empty')

        missing = [c for c in REQUIRED_COLUMNS if c not in headers]
        if missing:
            raise CommandError(f'missing columns: {missing}')
        idx = {h: i for i, h in enumerate(headers)}

        # Preload users for validation (modest scale; one query).
        users = {u.email.lower(): u for u in User.objects.all()}

        parsed, errors = [], []
        seen_external = set()
        seen_interviewee = {}  # user pk -> row number
        for n, row in enumerate(rows_iter, start=2):
            raw_ext = row[idx['ExternalID']]
            if not raw_ext:
                continue
            ext = str(raw_ext).strip()

            try:
                data = _parse_row(row, idx, users)
            except ValueError as e:
                errors.append(f'row {n} ({ext!r}): {e}')
                continue

            if ext in seen_external:
                errors.append(f'row {n}: duplicate ExternalID {ext!r}')
                continue
            seen_external.add(ext)

            interviewee = data['interviewee']
            if interviewee.pk in seen_interviewee:
                errors.append(
                    f'row {n} ({ext!r}): interviewee {interviewee.email} '
                    f'already used on row {seen_interviewee[interviewee.pk]}'
                )
                continue
            seen_interviewee[interviewee.pk] = n

            data['external_id'] = ext
            parsed.append(data)

        if errors:
            for e in errors:
                self.stderr.write(self.style.ERROR(e))
            raise CommandError(f'{len(errors)} row(s) failed; aborting')

        if dry_run:
            self.stdout.write(self.style.WARNING(
                f'dry-run: would upsert {len(parsed)} interviews'
            ))
            return

        created, updated = 0, 0
        try:
            with transaction.atomic():
                for data in parsed:
                    interview, was_created = Interview.objects.update_or_create(
                        external_id=data['external_id'],
                        defaults={
                            'title': data['title'],
                            'interviewee': data['interviewee'],
                        },
                    )
                    interview.panelists.set(data['panelists'])
                    created += was_created
                    updated += not was_created
        except IntegrityError as e:
            raise CommandError(f'database error during import (rolled back): {e}')

        self.stdout.write(self.style.SUCCESS(
            f'imported {len(parsed)} interviews ({created} created, {updated} updated)'
        ))


def _parse_row(row, idx, users):
    title = (row[idx['Title']] or '')
    title = str(title).strip()
    if not title:
        raise ValueError('missing Title')

    interviewee = _resolve(
        (row[idx['IntervieweeEmail']] or '').strip().lower(),
        users, User.Role.INTERVIEWEE, 'interviewee', required=True,
    )

    raw_panel = str(row[idx['PanelistEmails']] or '').strip()
    panelists, seen = [], set()
    for email in (e.strip().lower() for e in raw_panel.split(';') if e.strip()):
        if email in seen:
            continue
        seen.add(email)
        panelists.append(
            _resolve(email, users, User.Role.INTERVIEWER, 'panelist', required=True)
        )

    return {'title': title, 'interviewee': interviewee, 'panelists': panelists}


def _resolve(email, users, role, label, required):
    if not email:
        if required:
            raise ValueError(f'missing {label} email')
        return None
    user = users.get(email)
    if user is None:
        raise ValueError(f'unknown {label} email: {email}')
    if user.role != role:
        raise ValueError(f'{email} is not an {label} (role={user.role})')
    return user
