from datetime import datetime, time
from pathlib import Path
from zoneinfo import ZoneInfo

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from interviews.models import Timeslot

LONDON = ZoneInfo('Europe/London')
REQUIRED_COLUMNS = ('MeetingID', 'Date', 'DayLabel', 'StartTime', 'EndTime', 'JoinUrl', 'Status')


class Command(BaseCommand):
    help = 'Import (upsert) Timeslot rows from the bulk Teams meetings xlsx.'

    def add_arguments(self, parser):
        parser.add_argument('path', type=Path, help='Path to the source xlsx file')
        parser.add_argument('--sheet', default='Meetings', help='Worksheet name (default: Meetings)')
        parser.add_argument('--dry-run', action='store_true', help='Parse and validate but do not write')

    def handle(self, *args, path, sheet, dry_run, **_):
        import openpyxl

        if not path.exists():
            raise CommandError(f'file not found: {path}')

        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        if sheet not in wb.sheetnames:
            raise CommandError(f'sheet {sheet!r} not found; available: {wb.sheetnames}')

        ws = wb[sheet]
        rows_iter = ws.iter_rows(values_only=True)
        headers = next(rows_iter, None)
        if not headers:
            raise CommandError('sheet is empty')

        missing = [c for c in REQUIRED_COLUMNS if c not in headers]
        if missing:
            raise CommandError(f'missing columns: {missing}')
        idx = {h: i for i, h in enumerate(headers)}

        parsed, errors = [], []
        for n, row in enumerate(rows_iter, start=2):
            slot_ref = row[idx['MeetingID']]
            if not slot_ref:
                continue
            try:
                parsed.append(_parse_row(row, idx))
            except Exception as e:
                errors.append(f'row {n} ({slot_ref!r}): {e}')

        if errors:
            for e in errors:
                self.stderr.write(self.style.ERROR(e))
            raise CommandError(f'{len(errors)} row(s) failed; aborting')

        if dry_run:
            self.stdout.write(self.style.WARNING(
                f'dry-run: would upsert {len(parsed)} timeslots'
            ))
            return

        created, updated = 0, 0
        with transaction.atomic():
            for data in parsed:
                _, was_created = Timeslot.objects.update_or_create(
                    slot_ref=data['slot_ref'],
                    defaults=data,
                )
                created += was_created
                updated += not was_created

        self.stdout.write(self.style.SUCCESS(
            f'imported {len(parsed)} timeslots ({created} created, {updated} updated)'
        ))


def _parse_row(row, idx):
    slot_ref = str(row[idx['MeetingID']]).strip()
    date_val = row[idx['Date']]
    start_val = row[idx['StartTime']]
    end_val = row[idx['EndTime']]
    day_label = str(row[idx['DayLabel']]).strip()
    status = (row[idx['Status']] or '').strip()
    join_url = row[idx['JoinUrl']]

    try:
        suffix = int(slot_ref.rsplit('-', 1)[-1])
    except ValueError as e:
        raise ValueError(f'cannot derive room from MeetingID: {e}')
    room_label = f'Room {suffix}'

    return {
        'slot_ref': slot_ref,
        'day_label': day_label,
        'start_at': _combine(date_val, start_val),
        'end_at': _combine(date_val, end_val),
        'room_label': room_label,
        'teams_link': join_url if status == 'Created' and join_url else None,
    }


def _combine(date_val, time_val):
    if isinstance(time_val, time):
        t = time_val
    elif isinstance(time_val, str):
        h, m = time_val.split(':')[:2]
        t = time(int(h), int(m))
    else:
        raise ValueError(f'unsupported time value: {time_val!r}')

    if isinstance(date_val, datetime):
        d = date_val.date()
    elif hasattr(date_val, 'year'):
        d = date_val
    else:
        raise ValueError(f'unsupported date value: {date_val!r}')

    return datetime.combine(d, t).replace(tzinfo=LONDON)
